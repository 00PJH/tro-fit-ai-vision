"""
Tro-Fit SW Architecture Diagram Generator
=========================================
Design System: Meta-inspired light theme (white canvas, clean cards, thin borders)
Font: Pretendard ONLY
Slide Ratio: 16:9 (33.87 cm × 19.05 cm)
Tech Stack sources: STRICTLY from `2026_피우다프로젝트_WBS_로컬환경.md`
"""

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ─────────────────────────────────────────────
#  DESIGN SYSTEM TOKENS (Meta Commerce Light)
# ─────────────────────────────────────────────
C = {
    "canvas":           RGBColor(0xFF, 0xFF, 0xFF),   # White background
    "surface_soft":     RGBColor(0xF6, 0xF8, 0xFA),   # Light gray card bg
    "surface_card":     RGBColor(0xFF, 0xFF, 0xFF),   # Inner white card bg
    "hairline":         RGBColor(0xE1, 0xE4, 0xE6),   # 1px border
    "hairline_strong":  RGBColor(0xD1, 0xD5, 0xDA),   # Darker border
    "ink":              RGBColor(0x1F, 0x23, 0x28),   # Primary text
    "slate":            RGBColor(0x57, 0x60, 0x6A),   # Secondary text
    "steel":            RGBColor(0x8C, 0x95, 0x9F),   # Caption text
    "on_dark":          RGBColor(0xFF, 0xFF, 0xFF),   # White on colored bg
    "primary":          RGBColor(0x01, 0x66, 0xFF),   # Meta Cobalt
    "primary_soft":     RGBColor(0xE8, 0xF2, 0xFF),   # Soft Cobalt bg
    "tro_teal":         RGBColor(0x00, 0xB8, 0x82),   # Tro-Fit Teal
    "tro_soft":         RGBColor(0xE6, 0xF8, 0xF3),   # Soft Teal bg
}

FONT_NAME = "Pretendard"

SLIDE_W = Cm(33.87)
SLIDE_H = Cm(19.05)

MARGIN_L = Cm(2.2)
MARGIN_R = Cm(2.2)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R

# Positions
TOP_NAV_H    = Cm(1.10)
CHAPTER_TOP  = Cm(1.35)
DIVIDER_TOP  = Cm(2.00)
TITLE_TOP    = Cm(2.15)
SUBTITLE_TOP = Cm(3.20)
CONTENT_TOP  = Cm(4.00)
FOOTER_TOP   = Cm(17.85)

# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────

def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])

def set_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def draw_rect(slide, l, t, w, h, fill_color=None, line_color=None, line_pt=1.0):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_pt)
    else:
        shape.line.fill.background()
    return shape

def add_tb(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, spc=0.0, line_sp=None):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    if line_sp:
        p.line_spacing = line_sp
    run = p.add_run()
    run.text           = text
    run.font.name      = FONT_NAME
    run.font.size      = size
    run.font.color.rgb = color
    run.font.bold      = bold
    if spc:
        run._r.get_or_add_rPr().set("spc", str(int(spc * 100)))
    return box

def draw_chrome(slide, chapter: str, slide_num: str, source: str = ""):
    # Top nav
    draw_rect(slide, 0, 0, SLIDE_W, TOP_NAV_H, fill_color=C["canvas"])
    draw_rect(slide, 0, TOP_NAV_H, SLIDE_W, Pt(1), fill_color=C["hairline"])
    add_tb(slide, MARGIN_L, Pt(4), Cm(10), TOP_NAV_H, "TRO-FIT", Pt(11), C["ink"], bold=True, spc=1.5)
    add_tb(slide, SLIDE_W - Cm(3.8), Pt(4), Cm(3.5), TOP_NAV_H, slide_num, Pt(9), C["steel"], align=PP_ALIGN.RIGHT)

    # Chapter
    add_tb(slide, MARGIN_L, CHAPTER_TOP, CONTENT_W, Cm(0.6), chapter.upper(), Pt(8.5), C["primary"], bold=True, spc=1.8)
    draw_rect(slide, MARGIN_L - Cm(0.35), CHAPTER_TOP + Pt(2), Pt(4), Pt(11), fill_color=C["primary"])

    # Divider
    draw_rect(slide, MARGIN_L, DIVIDER_TOP, CONTENT_W, Pt(1), fill_color=C["hairline"])

    # Footer
    draw_rect(slide, MARGIN_L, FOOTER_TOP, CONTENT_W, Pt(1), fill_color=C["hairline"])
    if source:
        add_tb(slide, MARGIN_L, FOOTER_TOP + Pt(3), CONTENT_W * 0.8, Cm(0.55), source, Pt(7.5), C["steel"])

# ─────────────────────────────────────────────
#  BUILD SLIDE
# ─────────────────────────────────────────────

def build_sw_architecture_slide(prs: Presentation):
    slide = blank_slide(prs)
    set_bg(slide, C["canvas"])
    draw_chrome(slide, "System Architecture", "04b / 09", "출처: 2026 피우다프로젝트 WBS (로컬 환경 버전) 기술 세부 사항")

    # Titles
    add_tb(slide, MARGIN_L, TITLE_TOP, CONTENT_W, Cm(1.0), "Tro-Fit 소프트웨어 구성도 (SW 아키텍처)", Pt(26), C["ink"], bold=True)
    add_tb(slide, MARGIN_L, SUBTITLE_TOP, CONTENT_W, Cm(0.6), "로컬 Docker Compose 인프라 기반의 비동기식 온디바이스 AI 및 서비스 시스템 맵핑", Pt(12), C["slate"])

    # 3 Columns Layout (Client, Server, DB & Storage)
    col_w = Cm(8.6)
    col_h = Cm(9.6)
    gap = Cm(0.8)

    c1_l = MARGIN_L
    c2_l = c1_l + col_w + gap
    c3_l = c2_l + col_w + gap
    card_y = Cm(4.0)

    # ─────────────────────────────────────────
    # COLUMN 1: CLIENT TIER
    # ─────────────────────────────────────────
    draw_rect(slide, c1_l, card_y, col_w, col_h, fill_color=C["surface_soft"], line_color=C["hairline"], line_pt=1.0)
    add_tb(slide, c1_l + Pt(12), card_y + Pt(10), col_w - Pt(24), Pt(20), "CLIENT DEVICES", Pt(12), C["primary"], bold=True, spc=1.0)
    
    # 1.1 Mobile App Card
    m_y = card_y + Cm(0.9)
    draw_rect(slide, c1_l + Cm(0.3), m_y, col_w - Cm(0.6), Cm(3.8), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c1_l + Cm(0.5), m_y + Pt(6), col_w - Cm(1.0), Pt(16), "스마트폰 모바일 앱 (React Native / Expo)", Pt(10.5), C["ink"], bold=True)
    
    # Inner TFLite Engine Box
    tflite_y = m_y + Cm(0.9)
    draw_rect(slide, c1_l + Cm(0.5), tflite_y, col_w - Cm(1.0), Cm(2.4), fill_color=C["tro_soft"], line_color=C["tro_teal"], line_pt=1.0)
    add_tb(slide, c1_l + Cm(0.7), tflite_y + Pt(4), col_w - Cm(1.4), Pt(14), "On-device AI 엔진 (TFLite)", Pt(9.5), C["tro_teal"], bold=True)
    tflite_stack = (
        "• MediaPipe Pose (33개 3D 관절 추출)\n"
        "• RTMPose (모델 성능 비교/검증)\n"
        "• INT8 양자화 모델 경량화 적용"
    )
    add_tb(slide, c1_l + Cm(0.7), tflite_y + Pt(20), col_w - Cm(1.4), Cm(1.6), tflite_stack, Pt(8.5), C["slate"], line_sp=Pt(13))

    # 1.2 TV App / Mirroring Card
    tv_y = card_y + Cm(5.0)
    draw_rect(slide, c1_l + Cm(0.3), tv_y, col_w - Cm(0.6), Cm(3.8), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c1_l + Cm(0.5), tv_y + Pt(6), col_w - Cm(1.0), Pt(16), "스마트 TV 환경 (WebOS / Tizen)", Pt(10.5), C["ink"], bold=True)
    tv_stack = (
        "• Smart TV 전용 웹 애플리케이션 개발\n"
        "• HDMI 미러링 및 Chromecast SDK 연동\n"
        "• Lottie 웹 애니메이션 + HTML5 Audio 음악 재생\n"
        "• TV 리모컨 조작 인터랙션 연동"
    )
    add_tb(slide, c1_l + Cm(0.5), tv_y + Pt(22), col_w - Cm(1.0), Cm(3.0), tv_stack, Pt(8.5), C["slate"], line_sp=Pt(13))

    # ─────────────────────────────────────────
    # COLUMN 2: SERVER TIER (DOCKER COMPOSE)
    # ─────────────────────────────────────────
    draw_rect(slide, c2_l, card_y, col_w, col_h, fill_color=C["surface_soft"], line_color=C["hairline"], line_pt=1.0)
    add_tb(slide, c2_l + Pt(12), card_y + Pt(10), col_w - Pt(24), Pt(20), "LOCAL SERVER STACK (Docker)", Pt(12), C["tro_teal"], bold=True, spc=1.0)

    # 2.1 Backend API Server
    api_y = card_y + Cm(0.9)
    draw_rect(slide, c2_l + Cm(0.3), api_y, col_w - Cm(0.6), Cm(3.8), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c2_l + Cm(0.5), api_y + Pt(6), col_w - Cm(1.0), Pt(16), "API 서버 (FastAPI)", Pt(10.5), C["ink"], bold=True)
    api_stack = (
        "• Python RESTful API 서버\n"
        "• JWT 인증 및 사용자 보안 관리\n"
        "• WebSocket 기반 실시간 운동 모니터링 API\n"
        "• 카카오톡 알림톡 공유 연동 처리"
    )
    add_tb(slide, c2_l + Cm(0.5), api_y + Pt(22), col_w - Cm(1.0), Cm(3.0), api_stack, Pt(8.5), C["slate"], line_sp=Pt(13))

    # 2.2 Celery / Redis Async Engine
    celery_y = card_y + Cm(5.0)
    draw_rect(slide, c2_l + Cm(0.3), celery_y, col_w - Cm(0.6), Cm(3.8), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c2_l + Cm(0.5), celery_y + Pt(6), col_w - Cm(1.0), Pt(16), "비동기 분석 및 RAG 연동 엔진", Pt(10.5), C["ink"], bold=True)
    celery_stack = (
        "• Celery Worker + Redis Message Broker\n"
        "• LLM Fine-tuning 기반 안무 자동 처방\n"
        "• FFmpeg 활용 음원 & 비디오 미디어 인코딩 전처리\n"
        "• 시계열 이상 탐지 (낙상 위기 사전 예보 분석)"
    )
    add_tb(slide, c2_l + Cm(0.5), celery_y + Pt(22), col_w - Cm(1.0), Cm(3.0), celery_stack, Pt(8.5), C["slate"], line_sp=Pt(13))

    # ─────────────────────────────────────────
    # COLUMN 3: DATABASE & STORAGE TIER
    # ─────────────────────────────────────────
    draw_rect(slide, c3_l, card_y, col_w, col_h, fill_color=C["surface_soft"], line_color=C["hairline"], line_pt=1.0)
    add_tb(slide, c3_l + Pt(12), card_y + Pt(10), col_w - Pt(24), Pt(20), "DATABASE & STORAGE", Pt(12), C["ink"], bold=True, spc=1.0)

    # 3.1 PostgreSQL
    pg_y = card_y + Cm(0.9)
    draw_rect(slide, c3_l + Cm(0.3), pg_y, col_w - Cm(0.6), Cm(2.4), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c3_l + Cm(0.5), pg_y + Pt(6), col_w - Cm(1.0), Pt(16), "PostgreSQL (Docker RDB)", Pt(10.5), C["ink"], bold=True)
    add_tb(slide, c3_l + Cm(0.5), pg_y + Pt(22), col_w - Cm(1.0), Cm(1.4), "• 사용자 기본 정보 및 메타데이터\n• 주간 운동 통계 및 누적 ROM 시계열 기록 관리", Pt(8.5), C["slate"], line_sp=Pt(13))

    # 3.2 ChromaDB
    chroma_y = card_y + Cm(3.6)
    draw_rect(slide, c3_l + Cm(0.3), chroma_y, col_w - Cm(0.6), Cm(2.5), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c3_l + Cm(0.5), chroma_y + Pt(6), col_w - Cm(1.0), Pt(16), "ChromaDB (로컬 벡터 DB)", Pt(10.5), C["ink"], bold=True)
    add_tb(slide, c3_l + Cm(0.5), chroma_y + Pt(22), col_w - Cm(1.0), Cm(1.5), "• 36개 임상 진료 지침 및 가이던스 임베딩\n• LLM 안무 맵핑용 지식 데이터 및 유사도 검색", Pt(8.5), C["slate"], line_sp=Pt(13))

    # 3.3 MinIO
    minio_y = card_y + Cm(6.4)
    draw_rect(slide, c3_l + Cm(0.3), minio_y, col_w - Cm(0.6), Cm(2.4), fill_color=C["canvas"], line_color=C["hairline_strong"], line_pt=1.0)
    add_tb(slide, c3_l + Cm(0.5), minio_y + Pt(6), col_w - Cm(1.0), Pt(16), "MinIO (로컬 S3 호환 Object Storage)", Pt(10.5), C["ink"], bold=True)
    add_tb(slide, c3_l + Cm(0.5), minio_y + Pt(22), col_w - Cm(1.0), Cm(1.4), "• 생성된 트로트 안무 JSON 패킷 및 메타 리소스 보관\n• 로컬 S3 API 엔드포인트 연동", Pt(8.5), C["slate"], line_sp=Pt(13))

    # ─────────────────────────────────────────
    # DEVSECOPS & INFRASTRUCTURE ROW (BOTTOM)
    # ─────────────────────────────────────────
    infra_y = card_y + col_h + Cm(0.35)
    infra_h = Cm(1.8)
    draw_rect(slide, MARGIN_L, infra_y, CONTENT_W, infra_h, fill_color=C["primary_soft"], line_color=C["primary"], line_pt=1.0)
    add_tb(slide, MARGIN_L + Pt(14), infra_y + Pt(8), CONTENT_W - Pt(28), Pt(14), "DEVSECOPS & INFRASTRUCTURE TOOLS (로컬 개발 환경 검증)", Pt(9.5), C["primary"], bold=True, spc=1.0)

    # 4 Sub-blocks in DevSecOps
    infra_w = (CONTENT_W - Cm(0.9)) / 4
    infra_items = [
        ("CI/CD Pipeline", "GitHub Actions\n로컬 Docker Compose 빌드 자동화"),
        ("Monitoring", "Prometheus + Grafana\nDocker 멀티컨테이너 리소스 모니터링"),
        ("Performance Test", "Locust 활용\n로컬 API 서버 부하 테스트"),
        ("Security Audit", "OWASP ZAP 로컬 취약점 점검\nBandit 정적 분석")
    ]
    for ii, (ititle, idesc) in enumerate(infra_items):
        ix = MARGIN_L + ii * (infra_w + Cm(0.3)) + Pt(10)
        draw_rect(slide, MARGIN_L + ii * (infra_w + Cm(0.3)) + Cm(0.2), infra_y + Cm(0.6), infra_w - Cm(0.4), Cm(1.0), fill_color=C["canvas"], line_color=C["hairline"], line_pt=1.0)
        add_tb(slide, ix + Pt(6), infra_y + Cm(0.65), infra_w - Cm(0.6), Pt(12), ititle, Pt(9), C["ink"], bold=True)
        add_tb(slide, ix + Pt(6), infra_y + Cm(1.05), infra_w - Cm(0.6), Pt(18), idesc.replace('\n', '  |  '), Pt(7.5), C["slate"])

    # ─────────────────────────────────────────
    # CONNECTOR FLOW INDICATORS (SIMPLE & ROBUST)
    # ─────────────────────────────────────────
    # Client <-> Server
    c1_c2_x = c1_l + col_w + Cm(0.1)
    c1_c2_y = card_y + col_h / 2 - Cm(0.5)
    add_tb(slide, c1_c2_x, c1_c2_y, gap - Cm(0.2), Cm(1.0), "↔", Pt(20), C["primary"], bold=True, align=PP_ALIGN.CENTER)
    add_tb(slide, c1_c2_x, c1_c2_y + Cm(0.7), gap - Cm(0.2), Cm(0.5), "REST API\nJSON", Pt(7), C["steel"], align=PP_ALIGN.CENTER)

    # Server <-> DB
    c2_c3_x = c2_l + col_w + Cm(0.1)
    c2_c3_y = card_y + col_h / 2 - Cm(0.5)
    add_tb(slide, c2_c3_x, c2_c3_y, gap - Cm(0.2), Cm(1.0), "↔", Pt(20), C["tro_teal"], bold=True, align=PP_ALIGN.CENTER)
    add_tb(slide, c2_c3_x, c2_c3_y + Cm(0.7), gap - Cm(0.2), Cm(0.5), "SQL / Query\nS3 API", Pt(7), C["steel"], align=PP_ALIGN.CENTER)

    # Footer Stripe Accent
    draw_rect(slide, 0, SLIDE_H - Pt(4), SLIDE_W, Pt(4), fill_color=C["primary"])

if __name__ == "__main__":
    prs = new_prs()
    build_sw_architecture_slide(prs)
    out_file = "c:/workspace/trofit/middle_ppt/tro_fit_sw_architecture.pptx"
    prs.save(out_file)
    print(f"[OK] Saved SW Architecture slide to: {out_file}")
